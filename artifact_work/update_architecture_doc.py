from pathlib import Path
import shutil
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
FIGURES = Path(sys.argv[3])


def paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = Paragraph(new_p, paragraph._parent)
    if style:
        result.style = style
    if text:
        result.add_run(text)
    return result


def remove_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_caption(paragraph, text, italic=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.italic = italic


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def figure_after(anchor, image_name, caption, description, width=6.15):
    p_img = paragraph_after(anchor)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    shape = p_img.add_run().add_picture(str(FIGURES / image_name), width=Inches(width))
    set_alt_text(shape, caption, description)
    p_cap = paragraph_after(p_img)
    set_caption(p_cap, caption)
    p_src = paragraph_after(p_cap)
    set_caption(p_src, "Sumber: Hasil perancangan peneliti berdasarkan implementasi program", italic=True)
    p_src.paragraph_format.space_after = Pt(8)
    return p_src


def find_prefix(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"Paragraph not found: {prefix}")


def replace_text_prefix(doc, prefix, new_text):
    p = find_prefix(doc, prefix)
    clear_paragraph(p)
    p.add_run(new_text)
    return p


def set_cell_text(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(9)


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_module_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:color"), "D9D9D9")
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
        set_cell_width(row.cells[0], 2.35)
        set_cell_width(row.cells[1], 4.15)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "45515E" if row_index == 0 else ("F3F5F7" if row_index % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9)
                    if row_index == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)


def style_heading(p):
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

# Replace the two existing overview figures with corrected versions.
picture_paragraphs = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
if len(picture_paragraphs) < 2:
    raise RuntimeError("Expected at least two figure paragraphs")
for p, image_name, title, desc in [
    (picture_paragraphs[0], "figure_4_15_architecture_overview.png", "Gambar 4.15 Arsitektur sistem QA RAG pada lingkungan lokal", "Empat komponen logis sistem berada pada satu laptop lokal dan dikoordinasikan oleh FastAPI melalui Uvicorn."),
    (picture_paragraphs[1], "figure_4_16_question_sequence.png", "Gambar 4.16 Alur komunikasi satu pertanyaan pada arsitektur lokal", "Diagram urutan komunikasi Client, API Server, Retrieval and Knowledge Service, serta Inference Service."),
]:
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(FIGURES / image_name), width=Inches(6.15))
    set_alt_text(shape, title, desc)

# Correct statements that no longer match the current pipeline.
replace_text_prefix(doc, "Gambar 4.15 menunjukkan", "Gambar 4.15 menunjukkan bahwa API Server menjadi satu-satunya penghubung client dengan komponen pemrosesan. Panah dua arah antara Client dan API Server menunjukkan pengiriman permintaan serta pengembalian respons. Hubungan API Server dengan Retrieval and Knowledge Service digunakan untuk mengirim query dan memperoleh chunk sumber, sedangkan hubungan API Server dengan Inference Service digunakan untuk perluasan pertanyaan dan generasi jawaban.")
replace_text_prefix(doc, "6)  Inference Service ke API Server", "6)  Inference Service ke API Server. Ollama menjalankan model GGUF dan mengirimkan token secara bertahap. FastAPI meneruskan setiap token sebagai peristiwa SSE berstatus sementara. Setelah generasi selesai, API Server memeriksa jawaban terhadap sumber yang digunakan. Pemeriksaan meliputi keberadaan sitasi, validitas nomor sitasi, keberadaan angka jawaban pada sumber yang dirujuk, dan penggunaan angka target atau proyeksi pada klaim yang tidak diberi keterangan.")
replace_text_prefix(doc, "Pemeriksaan ini bersifat deterministik", "Pemeriksaan ini bersifat deterministik dan konservatif; hasil lolos tidak sama dengan jaminan kebenaran semantik. Jika ditemukan masalah, jawaban tetap ditampilkan dengan status needs_review dan catatan yang perlu diperiksa pengguna. Cuplikan sumber dipilih secara deterministik dari teks chunk berdasarkan istilah pertanyaan, sedangkan teks chunk lengkap tetap tersedia pada kartu sumber.")
replace_text_prefix(doc, "7)  API Server ke Client", "7)  API Server ke Client. Selama pemrosesan, FastAPI mengirimkan peristiwa queued, expanding, expanded, retrieving, reranked, sources, generating, token, answer, dan done. Heartbeat dikirim secara berkala ketika tidak ada peristiwa baru agar koneksi tetap aktif. Jika pengguna membatalkan atau menutup koneksi, sinyal pembatalan diteruskan kepada pekerjaan backend.")
replace_text_prefix(doc, "Peristiwa akhir memuat", "Peristiwa akhir memuat jawaban, status pemeriksaan, daftar sumber, publikasi terkait, peringatan, waktu setiap tahap, informasi cache, konfigurasi, metrik generasi, dan identitas permintaan. JavaScript mengganti jawaban sementara dengan jawaban akhir serta memperbarui panel proses, kartu sumber, dan informasi waktu tanpa memuat ulang halaman.")
replace_text_prefix(doc, "Ringkasan sumber dibuat", "Cuplikan sumber dibentuk oleh fungsi source_excerpt pada evidence.py. Fungsi ini memilih bagian teks chunk yang paling berkaitan dengan istilah pertanyaan tanpa memanggil model generatif. Frontend menampilkan cuplikan tersebut bersama teks sumber lengkap sehingga pengguna dapat menelusuri konteks asli yang dipakai sistem.")
replace_text_prefix(doc, "Bagian utama antarmuka terdiri", "Bagian utama antarmuka terdiri atas identitas aplikasi, tombol Cara pakai dan Cara kerja, kolom pertanyaan, pilihan model, pilihan gaya jawaban, indikator kesiapan backend, tombol Tanyakan, dan tombol Batalkan. Di bawah formulir tersedia panel proses yang menampilkan tahapan sistem dan rincian kandidat, area jawaban, peringatan, informasi waktu, kartu sumber, serta daftar publikasi terkait.")
replace_text_prefix(doc, "Saat pertanyaan dikirim", "Saat pertanyaan dikirim, frontend membuka koneksi SSE ke endpoint streaming dan menampilkan tahapan proses sesuai nama peristiwa yang diterima. Token yang datang ditampilkan sebagai jawaban sementara. Setelah peristiwa answer dan done diterima, antarmuka menampilkan jawaban akhir, status pemeriksaan, waktu proses, serta sumber yang digunakan.")
replace_text_prefix(doc, "Setiap kartu sumber menampilkan", "Setiap kartu sumber menampilkan judul publikasi, tahun terbit, rentang halaman, skor reranker, tautan publikasi BPS, cuplikan asli yang relevan, dan teks chunk lengkap. Penyajian tersebut membantu pengguna membedakan jawaban model dari bukti yang berasal langsung dari publikasi. Peringatan ditampilkan apabila sebagian sumber tidak masuk ke konteks atau jawaban memerlukan pemeriksaan lebih lanjut.")

# Update architecture summary table.
for row in doc.tables[0].rows[1:]:
    if row.cells[0].text.strip() == "Inference Service":
        set_cell_text(row.cells[2], "Menjalankan perluasan pertanyaan dan generasi jawaban menggunakan model GGUF.")

# Update backend module table to reflect the actual files.
module_rows = [
    ("main.py", "Inisialisasi FastAPI, pemuatan sumber daya, endpoint, antrean, SSE, pembatalan, health, readiness, dan penyajian frontend statis."),
    ("schemas.py", "Validasi struktur pertanyaan, model, top_k, penggunaan RAG, dan gaya jawaban."),
    ("question_validation.py", "Pemeriksaan isi pertanyaan dan pemetaan pesan kesalahan masukan."),
    ("pipeline.py", "Orkestrasi ekspansi, retrieval, reranking, konteks, generasi, pemeriksaan bukti, dan penyusunan payload akhir."),
    ("retrieval_service.py", "Pemuatan basis pengetahuan, embedding query, FAISS, BM25, RRF, serta metadata sumber."),
    ("reranker_service.py", "Pemuatan cross-encoder dan pengurutan ulang kandidat berdasarkan pertanyaan asli."),
    ("query_expansion_service.py", "Generasi, validasi, fallback, dan cache perluasan pertanyaan."),
    ("prompt_constructor.py", "Penyusunan pesan, format konteks, gaya jawaban, dan pembatasan ukuran konteks."),
    ("generation_service.py", "Komunikasi HTTP dengan Ollama, streaming token, status model, dan metrik generasi."),
    ("answer_stream.py", "Pembatasan panjang keluaran dan penyelesaian jawaban pada batas kalimat yang utuh."),
    ("evidence.py", "Pemeriksaan sitasi, angka, status target atau proyeksi, abstensi, dan pemilihan cuplikan sumber."),
    ("cache.py", "Cache dalam memori dengan kapasitas terbatas dan time-to-live."),
    ("config.py", "Konfigurasi basis pengetahuan, model, retrieval, Ollama, antrean, keamanan, dan gaya jawaban."),
    ("inference_errors.py", "Pemetaan kegagalan Ollama dan koneksi menjadi kode kesalahan aplikasi yang aman."),
]
table = doc.tables[1]
while len(table.rows) > 1:
    tr = table.rows[-1]._tr
    tr.getparent().remove(tr)
for mod, responsibility in module_rows:
    cells = table.add_row().cells
    set_cell_text(cells[0], mod)
    set_cell_text(cells[1], responsibility)
format_module_table(table)

# Add detailed backend figure after the opening backend paragraph.
backend_intro = find_prefix(doc, "Backend dikembangkan secara modular")
anchor = figure_after(
    backend_intro,
    "figure_4_17_backend_detail.png",
    "Gambar 4.17 Rincian komponen API Server dan backend",
    "Diagram modul backend mulai dari Uvicorn, FastAPI, validasi dan antrean, orkestrasi pipeline, sampai pemeriksaan keluaran.",
)
extra = paragraph_after(anchor, "Gambar 4.17 memperlihatkan bahwa Uvicorn hanya menjalankan aplikasi ASGI, sedangkan FastAPI pada main.py menangani lapisan HTTP dan siklus hidup sistem. Permintaan yang telah divalidasi masuk ke JobRunner, kemudian pipeline.py mengoordinasikan modul ekspansi, retrieval, penyusunan prompt, komunikasi dengan Ollama, dan pemeriksaan keluaran. Modul konfigurasi, cache, serta penanganan kesalahan mendukung seluruh tahapan tanpa menjadi jalur pemrosesan yang berdiri sendiri.")

# Insert dedicated retrieval and inference sections before the frontend section.
frontend_heading = find_prefix(doc, "4.3.4 Pembangunan Antarmuka Pengguna")
anchor = frontend_heading._p.getprevious()
anchor_p = Paragraph(anchor, frontend_heading._parent)

retrieval_heading = paragraph_after(anchor_p, "4.3.4 Implementasi Retrieval and Knowledge Service", "Heading 3")
style_heading(retrieval_heading)
retrieval_heading.paragraph_format.page_break_before = True
p = paragraph_after(retrieval_heading, "Retrieval and Knowledge Service menggabungkan proses pencarian dengan artefak pengetahuan lokal. Pada tahap startup, faiss_index.bin dan chunk_metadata.jsonl dimuat, konsistensi identitas vektor diperiksa, model embedding disiapkan, serta indeks BM25 dibentuk dari chunk yang bukan daftar isi.")
p = paragraph_after(p, "Ketika pertanyaan diproses, pertanyaan asli dan perluasan yang valid digunakan pada dua jalur pencarian. Jalur dense membentuk embedding dengan Nomic Embed Text v1.5 dan mencari maksimal 20 hasil melalui FAISS. Jalur sparse menghitung kecocokan kata menggunakan BM25 dan mengambil maksimal 20 hasil dengan skor positif. Seluruh daftar kemudian digabungkan menggunakan reciprocal rank fusion dengan konstanta 60 dan dibatasi menjadi maksimal 20 kandidat unik.")
p = paragraph_after(p, "Cross-encoder madebyaris/rerank-indonesia menilai ulang setiap kandidat menggunakan pertanyaan asli. Kandidat yang memenuhi ambang skor kemudian diperiksa oleh fit_context. Sistem mengambil maksimal lima chunk utuh yang masih muat bersama instruksi dan ruang keluaran pada konteks 8.192 token. Retrieval Service mengembalikan teks chunk, skor, dan metadata sumber kepada pipeline; layanan ini tidak mengirimkan data langsung kepada Ollama.")
anchor = figure_after(p, "figure_4_18_retrieval_detail.png", "Gambar 4.18 Alur Retrieval and Knowledge Service", "Pencarian hibrida melalui Nomic Embed, FAISS, BM25, RRF, cross-encoder reranker, dan pembatasan konteks dengan domain knowledge lokal.")

inference_heading = paragraph_after(anchor, "4.3.5 Implementasi Inference Service dengan Ollama", "Heading 3")
style_heading(inference_heading)
inference_heading.paragraph_format.page_break_before = True
p = paragraph_after(inference_heading, "Inference Service dijalankan oleh Ollama pada alamat lokal http://127.0.0.1:11434. FastAPI berkomunikasi dengan layanan tersebut melalui generation_service.py menggunakan endpoint /api/chat untuk inferensi dan /api/tags untuk pemeriksaan model. Pemisahan ini bersifat logis dan proses; FastAPI, Ollama, model, dan basis pengetahuan tetap berada pada laptop lokal yang sama.")
p = paragraph_after(p, "Ollama digunakan pada dua tahap. Pertama, model menghasilkan satu perluasan query yang kemudian divalidasi oleh backend. Kedua, model menghasilkan jawaban dari prompt yang berisi instruksi, pertanyaan, serta chunk sumber terpilih. Pada tahap jawaban, Ollama mengirim token dan metrik secara bertahap dalam aliran JSON. generation_service.py membaca aliran tersebut, sedangkan pipeline.py meneruskannya kepada frontend melalui SSE.")
p = paragraph_after(p, "Model yang terdaftar terdiri atas Llama 3.2 3B dan Gemma 2 2B dalam versi dasar dan fine-tuned dengan kuantisasi Q4_K_M. Konfigurasi bawaan menggunakan temperature 0, seed 42, panjang konteks 8.192 token, num_batch 128, keep_alive 15 menit, dan num_gpu -1 sehingga penempatan CPU atau GPU ditentukan oleh Ollama. Batas keluaran adalah 150 token untuk gaya ringkas dan 384 token untuk gaya detail.")
anchor = figure_after(p, "figure_4_19_inference_detail.png", "Gambar 4.19 Alur Inference Service berbasis Ollama", "Komunikasi FastAPI dengan Ollama lokal untuk pemeriksaan model, ekspansi pertanyaan, generasi jawaban, streaming token, dan metrik inferensi.")

# Renumber and expand the frontend section, then add its diagram.
clear_paragraph(frontend_heading)
frontend_heading.add_run("4.3.6 Pembangunan Antarmuka Pengguna")
style_heading(frontend_heading)
frontend_heading.paragraph_format.page_break_before = True
frontend_intro = find_prefix(doc, "Antarmuka pengguna pada tahap ini")
figure_anchor = figure_after(frontend_intro, "figure_4_20_frontend_detail.png", "Gambar 4.20 Rincian komponen dan alur frontend", "Berkas frontend, komponen antarmuka, pemeriksaan health, validasi masukan, permintaan SSE, dan penyajian jawaban serta sumber.")
paragraph_after(figure_anchor, "Gambar 4.20 memisahkan berkas pembentuk tampilan dari logika interaksi. index.html dan styles.css membentuk struktur serta tata letak, sedangkan app.js mengelola pemeriksaan health, pengiriman permintaan, konsumsi SSE, pembatalan, dan pembaruan elemen halaman. question-validation.js dan question-rules.json memberikan umpan balik awal di browser, tetapi validasi FastAPI tetap menjadi pemeriksaan yang menentukan apakah permintaan diterima.")

limit_heading = find_prefix(doc, "4.3.5 Batas Implementasi pada Lingkungan Lokal")
clear_paragraph(limit_heading)
limit_heading.add_run("4.3.7 Batas Implementasi pada Lingkungan Lokal")
style_heading(limit_heading)

# Normalize caption/source spacing and heading colors.
for p in doc.paragraphs:
    if (p.text.startswith("Gambar 4.") and len(p.text) < 100) or p.text.startswith("Tabel 4."):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
    elif p.text.startswith(("Gambar 4.15 menunjukkan", "Gambar 4.17 memperlihatkan", "Gambar 4.20 memisahkan")):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if p.style and p.style.name.startswith("Heading"):
        style_heading(p)
    if p.text.startswith("Tabel 4."):
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)

doc.save(OUTPUT)
print(OUTPUT)
